"""Styled, local document generation with format-specific validation.

Agent OTG — DWE Team
All output files are saved to ~/Downloads/AgentOTG/ so they are immediately
accessible from the user's Downloads folder.
"""
from __future__ import annotations

import csv
import datetime
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch, cm
from reportlab.platypus import (
    Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle,
    HRFlowable, KeepTogether,
)


# ── Output directory: ~/Downloads/AgentOTG/ ───────────────────────────────────
OUTPUT_DIR = Path(os.environ.get("USERPROFILE", Path.home())) / "Downloads" / "AgentOTG"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
LOGO_PATH = Path(__file__).resolve().parent / "logo" / "otglogo.png"

SUPPORTED_FORMATS = {"pdf", "docx", "txt", "md", "pptx", "xlsx", "csv", "json"}

# ── Brand palette ─────────────────────────────────────────────────────────────
_TEAL      = "087E70"   # primary accent
_TEAL_DARK = "065F54"   # darker variant for headers
_TEAL_LIGHT= "E8F5F3"   # very light teal for alternating rows
_INK       = "18212B"   # body text
_PALE      = "EAF3F0"   # light background panels
_SLATE     = "4A5568"   # secondary text / footer
_WHITE     = "FFFFFF"
_WARN      = "C05621"   # callout warning color
_NOTE      = "2B6CB0"   # callout note color

_TODAY = datetime.date.today().strftime("%B %d, %Y")


@dataclass(frozen=True)
class Artifact:
    path: Path
    media_type: str
    size_bytes: int

    def as_dict(self) -> dict[str, Any]:
        return {
            "filename": self.path.name,
            "path": str(self.path),
            "media_type": self.media_type,
            "size_bytes": self.size_bytes,
        }


def _safe_filename(filename: str, extension: str) -> Path:
    stem = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(str(filename or "output")).name).strip(" .") or "output"
    if Path(stem).suffix.lower() != f".{extension}":
        stem = f"{Path(stem).stem}.{extension}"
    return OUTPUT_DIR / stem


def _paragraphs(content: str) -> list[str]:
    return [item.strip() for item in re.split(r"\n\s*\n", content or "") if item.strip()]


def _rows(content: str) -> list[list[str]]:
    """Parse tabular content from CSV, TSV, or Markdown tables into normalized rows.
    Strips markdown code fences (```csv) and conversational preambles safely.
    """
    if not content or not str(content).strip():
        return [["Content"], ["No content supplied"]]

    text = str(content).strip()

    # 1. If fenced code block is present, extract content inside fences
    fence_pattern = re.compile(r"```(?:csv|tsv|excel|markdown|table)?\s*\n(.*?)\n```", re.DOTALL | re.IGNORECASE)
    fence_match = fence_pattern.search(text)
    if fence_match:
        text = fence_match.group(1).strip()
    else:
        # Strip any leading or trailing fence line
        text = re.sub(r"^```[a-zA-Z0-9_-]*\s*\n?", "", text)
        text = re.sub(r"\n?```\s*$", "", text).strip()

    raw_lines = text.splitlines()
    cleaned_lines = []

    for line in raw_lines:
        s = line.strip()
        if not s or s.startswith("```"):
            continue
        cleaned_lines.append(s)

    if not cleaned_lines:
        return [["Content"], ["No content supplied"]]

    # 2. Check if content is a Markdown table (lines starting with | and containing |)
    pipe_lines = [l for l in cleaned_lines if l.startswith("|") and "|" in l[1:]]
    if len(pipe_lines) >= 2:
        table_rows = []
        for pl in pipe_lines:
            # Skip separator rows like |---|---| or |:---|:---|
            if re.match(r"^\|[-| :]+\|$", pl):
                continue
            cells = [c.strip() for c in pl.strip("|").split("|")]
            if any(cells):
                table_rows.append(cells)
        if len(table_rows) >= 1:
            max_cols = max(len(r) for r in table_rows)
            return [r + [""] * (max_cols - len(r)) for r in table_rows]

    # 3. Check CSV / TSV format
    data_lines = []
    for l in cleaned_lines:
        # Ignore conversational preambles/postambles
        if re.match(r"^(?:here is|here's|below is|the following|sure|certainly|table:|data:|hope this helps|let me know|regards|thank)", l, re.I):
            continue
        data_lines.append(l)

    if not data_lines:
        data_lines = cleaned_lines

    # Delimiter detection: tab vs comma
    sample = "\n".join(data_lines[:5])
    delimiter = "\t" if "\t" in sample and sample.count("\t") >= sample.count(",") else ","

    try:
        reader = csv.reader(data_lines, delimiter=delimiter)
        rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    except Exception:
        rows = []

    if rows:
        max_cols = max(len(r) for r in rows)
        return [r + [""] * (max_cols - len(r)) for r in rows]

    return [["Content"], [content.strip() or "No content supplied"]]


def _parse_cell_value(val: Any) -> Any:
    """Parse cell string into native Python int, float, or bool for Excel.
    Preserves strings with leading zeroes (e.g. '0123') and alphanumeric IDs ('STU-001').
    """
    if not isinstance(val, str):
        return val
    s = val.strip()
    if not s:
        return ""
    # Booleans
    if s.lower() == "true":
        return True
    if s.lower() == "false":
        return False
    # Integers: allow negative, preserve leading zero codes (e.g. "01234") unless it's "0"
    if re.fullmatch(r"-?\d+", s):
        if len(s) > 1 and s.startswith("0"):
            return s
        if len(s) > 2 and s.startswith("-0"):
            return s
        try:
            return int(s)
        except ValueError:
            return s
    # Floats: e.g. "12.34", "-0.5"
    if re.fullmatch(r"-?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            return s
    # Currency values: e.g. "$1,250.00", "$45", "€100"
    m_curr = re.fullmatch(r"[\$€£¥]\s*(-?[\d,]+(?:\.\d+)?)", s)
    if m_curr:
        raw_num = m_curr.group(1).replace(",", "")
        try:
            return float(raw_num) if "." in raw_num else int(raw_num)
        except ValueError:
            return s
    return s


def _clean_topic_from_question(question: str) -> str:
    """Extract a clean, concise, 2-5 word professional topic title from a question prompt."""
    if not question:
        return "Generated Document"
    q = question.strip()

    # 0. Remove conversational intention wrappers like:
    # "I want to watch best tmkoc episodes ever. so give me the word file showing all the best episodes of tmkoc ever"
    # "I need a report on solar energy. Please generate a word file"
    q = re.sub(r"^\s*(?:i\s+(?:want|need|would\s+like)\s+(?:to\s+)?.*?(?:[.;,]\s*|\bso\b\s+|\bthen\b\s+))", "", q, flags=re.I)
    q = re.sub(r"^\s*(?:i\s+(?:want|need|would\s+like)(?:\s+to)?\s+)", "", q, flags=re.I)
    q = re.sub(r"^\s*(?:so\s+)?(?:give|show|make|create|write|generate|produce|draft|prepare|provide)\s+(?:me\s+)?(?:the|a|an)?\s*", "", q, flags=re.I)

    # 1. Remove leading conversational phrases
    q = re.sub(r"^\s*(?:please\s+)?(?:can you\s+)?(?:give me|show me|write|create|generate|make|build|export|produce|draft|prepare|provide)\s+", "", q, flags=re.I)

    # 2. Remove file format nouns and dataset/report descriptors at the start
    q = re.sub(r"^\s*(?:(?:a|an|the)\s+)?(?:pdf|docx|doc|word(?:\s+document|\s+doc|\s+file|\s+report)?|excel(?:\s+sheet|\s+file)?|xlsx|"
               r"spreadsheet|powerpoint|pptx|ppt|presentation|report|file|script|code|program|dataset|data)\s*", "", q, flags=re.I)

    # 3. 'showing', 'displaying', 'listing (out)?', 'containing', 'highlighting', 'explaining', 'describing', 'discussing', 'covering'
    q = re.sub(r"^\s*(?:showing|displaying|listing\s+(?:out\s+)?|containing|highlighting|explaining|describing|discussing|covering)\s+(?:all\s+(?:the\s+)?)?", "", q, flags=re.I)

    # 4. Remove leading prepositions and topic clauses ('about', 'on', 'for', 'covering', 'regarding', 'of', 'with', 'topic of/on')
    q = re.sub(r"^\s*(?:about|on|for|covering|regarding|of|related\s+to|titled|named|called)\s+", "", q, flags=re.I)
    q = re.sub(r"^\s*(?:the\s+)?topic\s+(?:of|on)?\s*", "", q, flags=re.I)

    # 5. Remove second wave of file descriptors if chained (e.g. "report on", "dataset of", "python script for")
    q = re.sub(r"^\s*(?:(?:a|an|the)\s+)?(?:report|document|presentation|spreadsheet|dataset|file|sheet|code|script)\s+(?:about|on|for|covering|regarding|of|to)\s+", "", q, flags=re.I)
    q = re.sub(r"^\s*(?:(?:a|an|the)\s+)?(?:python|javascript|java|c\+\+|sql|bash|shell)?\s*(?:script|code|program|function|class)\s+(?:for|to|of|about)\s+", "", q, flags=re.I)

    # 6. Remove trailing file creation clauses (e.g. "and also generate the pdf", "and put it in a word file")
    q = re.sub(
        r"(?:,?\s*(?:and\s+)?(?:also\s+)?)?(?:generate|create|make|build|export|produce|save|put|write|draft|prepare|provide)"
        r"(?:\s+it)?(?:\s+as|\s+in|\s+to)?(?:\s+an?|\s+the)?\s*"
        r"(?:pdf|docx|doc|word(?:\s+document|\s+doc|\s+file)?|excel(?:\s+sheet|\s+file)?|xlsx|spreadsheet|"
        r"powerpoint|pptx|ppt|presentation|report|file)\b.*$",
        "", q, flags=re.I,
    )

    # 7. Remove trailing constraint instructions ("with 5 key sections", "in 10 points", "with 3 bullets", "and 5 sample rows")
    q = re.sub(
        r"(?:,?\s*)?(?:with|containing|including|having|in)\s+\d+\s+(?:key\s+)?(?:sections?|points?|bullets?|highlights?|rows?|columns?|pages?|sample\s+rows?).*$",
        "", q, flags=re.I,
    )
    q = re.sub(r"(?:,?\s*)?(?:with|containing|including)\s+(?:code\s+snippets?|examples?|sample\s+data|tables?|best\s+practices?).*$", "", q, flags=re.I)
    # Remove trailing format specifications (e.g. "in docx and pdf and ppt", "as a pdf", "in word format")
    q = re.sub(
        r"(?:,?\s*)?(?:in|as|into|to)\s+(?:(?:an?|the)\s+)?(?:pdf|docx?|doc|word(?:\s+document|\s+doc|\s+file)?|excel(?:\s+sheet|\s+file)?|xlsx|spreadsheet|powerpoint|pptx?|ppt|presentation|markdown|text|json|csv)(?:\s*(?:and|,|/|\+)\s*(?:an?\s+)?(?:pdf|docx?|doc|word(?:\s+document|\s+doc|\s+file)?|excel(?:\s+sheet|\s+file)?|xlsx|spreadsheet|powerpoint|pptx?|ppt|presentation|markdown|text|json|csv))*\s*(?:format|file|document|deck)?\s*$",
        "", q, flags=re.I
    )
    # Remove "in python", "in excel", "in word" if at the very end
    q = re.sub(r"\s+(?:in|using)\s+(?:python|excel|word|powerpoint)\s*$", "", q, flags=re.I)
    q = re.sub(r"\s+\b(?:ever|of\s+all\s+time)\b\s*$", "", q, flags=re.I)

    # 8. Strip surrounding punctuation
    q = q.strip(" ,.;:-_")
    q = re.sub(r"^\s*(?:about|on|for|covering|regarding|of)\s+", "", q, flags=re.I).strip(" ,.;:-_")

    if not q or len(q) < 2:
        return "Generated Document"

    # Title case formatting for clean presentation
    words = q.split()
    if len(words) <= 6:
        q = " ".join(w.capitalize() if w.lower() not in ("a", "an", "the", "in", "on", "for", "and", "of", "to", "with", "vs") else w.lower() for w in words)
        q = q[0].upper() + q[1:] if q else q

    return q


def derive_clean_title(question: str, content: str = "") -> tuple[str, str, str]:
    """
    Derives a clean document title, sanitized content (without duplicated title heading),
    and a filesystem-safe slug for the filename.
    Prevents raw prompt instructions (e.g. 'Give me the code to...') and generic section
    headings (e.g. 'Introduction', 'Overview', 'Summary') from becoming the main heading or filename.
    """
    GENERIC_TITLES = {
        "introduction", "intro", "overview", "summary", "executive summary",
        "table of contents", "background", "abstract", "preface",
        "contents", "agenda", "slide 1", "slide", "title", "document",
        "output", "report", "presentation", "presentation brief",
        "brief", "notes", "section 1", "chapter 1",
    }

    # 0. Check if user explicitly asked for a title in quotes, e.g. titled "Q3 Report"
    quoted_title_match = re.search(r'(?:title|titled|named|called)\s+["\'"]([^"\']+)["\'"]', question or "", re.I)
    if quoted_title_match:
        explicit_title = quoted_title_match.group(1).strip()
        slug = re.sub(r'[^a-zA-Z0-9]+', '_', explicit_title.lower()).strip('_')
        slug_parts = [p for p in slug.split('_') if p][:5]
        safe_name = '_'.join(slug_parts) if slug_parts else 'document'
        return explicit_title, content, safe_name

    # 1. Check content for leading heading (# Heading or Title: ...)
    content_lines = (content or "").splitlines()
    title_from_content = ""
    new_content_lines = []
    found_heading = False

    for i, line in enumerate(content_lines):
        stripped = line.strip()
        if not found_heading and i < 5 and (stripped.startswith("#") or stripped.lower().startswith("title:")):
            if stripped.startswith("#"):
                raw_h = re.sub(r"^#+\s*", "", stripped).strip()
            else:
                raw_h = re.sub(r"^title:\s*", "", stripped, flags=re.I).strip()
            raw_h = re.sub(r"[\*`_]", "", raw_h).strip()

            # Check if this heading is a raw prompt command rather than a document title
            is_prompt_command = bool(re.match(
                r"^\s*(?:please\s+)?(?:can you\s+)?(?:give me|write|create|generate|make|build|export|produce)\s+(?:an?\s+)?(?:pdf|word|docx?|excel|xlsx|pptx?|powerpoint|spreadsheet|report|file)\b",
                raw_h, re.I
            ))

            # Generic headings like 'Introduction', 'Overview', 'Summary' are section titles, not document titles
            clean_raw_h = raw_h.lower().strip(" :.-_")
            is_generic_heading = clean_raw_h in GENERIC_TITLES or bool(re.match(
                r"^\s*(?:introduction|intro|overview|summary|executive\s+summary|background|abstract|preface|"
                r"table\s+of\s+contents|contents|agenda|slide\s*\d*|section\s*\d*|chapter\s*\d*|title|document|report|brief)\s*[:\-–—]?\s*$",
                raw_h.strip(), re.I
            ))

            if is_generic_heading or stripped.startswith("##"):
                # The document begins directly with section headings (e.g. ## Introduction).
                # Do NOT scan later sections as the document title; preserve all lines and use topic from prompt.
                new_content_lines.extend(content_lines[i:])
                break

            if len(raw_h) > 2 and not is_prompt_command:
                title_from_content = raw_h
                found_heading = True
                continue
        new_content_lines.append(line)

    if found_heading and title_from_content:
        clean_title = re.sub(
            r"^\s*(?:please\s+)?(?:can you\s+)?(?:give me|write|create|generate|make|build|export|produce)\s+"
            r"(?:(?:a|an|the)\s+)?(?:pdf|docx|word(?:\s+document|\s+doc)?|excel(?:\s+sheet)?|xlsx|"
            r"spreadsheet|powerpoint|pptx|ppt|presentation|report|file)\s*"
            r"(?:about|on|for|covering|regarding)?\s*",
            "", title_from_content, flags=re.I,
        ).strip(" ,.;:-") or title_from_content
        cleaned_content = "\n".join(new_content_lines).strip()
    else:
        cleaned_content = content
        clean_title = _clean_topic_from_question(question)

    # If clean_title is empty or in generic titles, fallback to topic from question
    if not clean_title or clean_title.lower().strip(" :.-_") in GENERIC_TITLES:
        clean_title = _clean_topic_from_question(question)
    if not clean_title or clean_title.lower().strip(" :.-_") in GENERIC_TITLES:
        clean_title = "Executive Report"

    clean_title = re.sub(r"^#+\s*", "", clean_title).strip()
    clean_title = re.sub(r"\s+", " ", clean_title).strip()
    clean_title = clean_title[:110].rstrip(" ,.;:-")

    # Generate safe_name filename slug strictly from clean_title
    slug_raw = re.sub(r"[^a-zA-Z0-9]+", "_", clean_title.lower()).strip("_")
    meta_words = {"a", "an", "the", "in", "on", "for", "and", "of", "to", "also", "with", "me", "give", "code", "script", "generate", "create", "file", "document", "report", "pdf", "docx", "word", "excel", "xlsx", "pptx", "powerpoint", "put", "it", "as"}
    slug_parts = [p for p in slug_raw.split("_") if p and p not in meta_words][:5]
    if not slug_parts:
        slug_parts = [p for p in slug_raw.split("_") if p][:5]
    safe_name = "_".join(slug_parts) if slug_parts else "document"

    return clean_title, cleaned_content, safe_name



def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ══════════════════════════════════════════════════════════════════════════════
# ADVANCED PDF GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def _make_styles() -> dict:
    """Create a complete style registry for advanced PDF layout."""
    base = getSampleStyleSheet()
    W, H = A4

    styles = {}

    styles["cover_title"] = ParagraphStyle(
        "CoverTitle",
        fontName="Helvetica-Bold",
        fontSize=32,
        leading=38,
        textColor=colors.HexColor(f"#FFFFFF"),
        spaceAfter=10,
        alignment=TA_LEFT,
    )
    styles["cover_subtitle"] = ParagraphStyle(
        "CoverSubtitle",
        fontName="Helvetica",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#CCE8E4"),
        spaceAfter=6,
        alignment=TA_LEFT,
    )
    styles["cover_meta"] = ParagraphStyle(
        "CoverMeta",
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#A0D4CE"),
        alignment=TA_LEFT,
    )
    styles["h1"] = ParagraphStyle(
        "AgentH1",
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor(f"#{_TEAL_DARK}"),
        spaceBefore=18,
        spaceAfter=8,
        borderPadding=(0, 0, 4, 0),
    )
    styles["h2"] = ParagraphStyle(
        "AgentH2",
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor(f"#{_TEAL}"),
        spaceBefore=14,
        spaceAfter=6,
    )
    styles["h3"] = ParagraphStyle(
        "AgentH3",
        fontName="Helvetica-BoldOblique",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor(f"#{_INK}"),
        spaceBefore=10,
        spaceAfter=4,
    )
    styles["body"] = ParagraphStyle(
        "AgentBody",
        fontName="Helvetica",
        fontSize=10.5,
        leading=15.5,
        textColor=colors.HexColor(f"#{_INK}"),
        spaceAfter=7,
        alignment=TA_JUSTIFY,
    )
    styles["bullet"] = ParagraphStyle(
        "AgentBullet",
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor(f"#{_INK}"),
        leftIndent=18,
        firstLineIndent=-12,
        spaceAfter=4,
    )
    styles["bullet2"] = ParagraphStyle(
        "AgentBullet2",
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor(f"#{_SLATE}"),
        leftIndent=34,
        firstLineIndent=-12,
        spaceAfter=3,
    )
    styles["numbered"] = ParagraphStyle(
        "AgentNumbered",
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor(f"#{_INK}"),
        leftIndent=22,
        firstLineIndent=-16,
        spaceAfter=4,
    )
    styles["code"] = ParagraphStyle(
        "AgentCode",
        fontName="Courier",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#2D3748"),
        backColor=colors.HexColor("#F7FAFC"),
        borderColor=colors.HexColor("#CBD5E0"),
        borderWidth=0.75,
        borderPadding=(8, 10, 8, 10),
        spaceBefore=6,
        spaceAfter=10,
    )
    styles["callout_note"] = ParagraphStyle(
        "AgentCalloutNote",
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor(f"#{_NOTE}"),
        backColor=colors.HexColor("#EBF8FF"),
        borderColor=colors.HexColor(f"#{_NOTE}"),
        borderWidth=1,
        borderPadding=(8, 10, 8, 10),
        spaceBefore=6,
        spaceAfter=8,
        leftIndent=4,
    )
    styles["callout_warn"] = ParagraphStyle(
        "AgentCalloutWarn",
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor(f"#{_WARN}"),
        backColor=colors.HexColor("#FFFAF0"),
        borderColor=colors.HexColor(f"#{_WARN}"),
        borderWidth=1,
        borderPadding=(8, 10, 8, 10),
        spaceBefore=6,
        spaceAfter=8,
        leftIndent=4,
    )
    styles["footer"] = ParagraphStyle(
        "AgentFooter",
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor(f"#{_SLATE}"),
        alignment=TA_CENTER,
    )
    return styles


def _pdf_cover_page(canvas, document, title: str) -> None:
    """Draw the executive cover page header band with white title text on teal background."""
    W, H = A4
    canvas.saveState()
    # Full-width teal cover band (top ~2.4 inches of first page)
    band_h = 2.45 * inch
    canvas.setFillColor(colors.HexColor(f"#{_TEAL_DARK}"))
    canvas.rect(0, H - band_h, W, band_h, stroke=0, fill=1)

    # Decorative accent stripe
    canvas.setFillColor(colors.HexColor(f"#{_TEAL}"))
    canvas.rect(0, H - band_h - 4, W, 4, stroke=0, fill=1)

    # Left sidebar accent (brighter teal)
    canvas.setFillColor(colors.HexColor("#04A896"))
    canvas.rect(0, H - band_h, 7, band_h, stroke=0, fill=1)

    # Agent OTG logo text (top-left)
    canvas.setFillColor(colors.HexColor("#A0D4CE"))
    canvas.setFont("Helvetica-Bold", 8.5)
    canvas.drawString(0.72 * inch, H - 0.42 * inch, "★  AGENT OTG  |  DWE TEAM  •  EXECUTIVE INTELLIGENCE")
    if LOGO_PATH.is_file():
        # The supplied white, transparent logo is deliberately placed on the
        # dark cover band so it remains crisp and legible.
        canvas.drawImage(str(LOGO_PATH), W - 1.38 * inch, H - 1.42 * inch,
                         width=0.72 * inch, height=0.72 * inch, mask="auto", preserveAspectRatio=True)

    # Divider line in header
    canvas.setStrokeColor(colors.HexColor("#04A896"))
    canvas.setLineWidth(0.6)
    canvas.line(0.72 * inch, H - 0.55 * inch, W - 0.72 * inch, H - 0.55 * inch)

    # Clean executive title rendered inside the teal band
    canvas.setFont("Helvetica-Bold", 20)
    canvas.setFillColor(colors.white)

    # Smart wrapping for title
    words = title.split()
    t_lines = []
    cur = []
    for w in words:
        cur.append(w)
        if len(" ".join(cur)) > 42:
            t_lines.append(" ".join(cur[:-1]))
            cur = [w]
    if cur:
        t_lines.append(" ".join(cur))

    y_pos = H - 0.95 * inch
    for line in t_lines[:2]:
        canvas.drawString(0.72 * inch, y_pos, line)
        y_pos -= 0.32 * inch

    # Subtitle metadata bar inside cover band
    canvas.setFont("Helvetica", 8.5)
    canvas.setFillColor(colors.HexColor("#CCE8E4"))
    meta_str = f"Date: {_TODAY}   |   Classification: Operational Report   |   Status: Verified & On-Premise"
    canvas.drawString(0.72 * inch, H - 2.05 * inch, meta_str)

    canvas.restoreState()


def _pdf_footer_only(canvas, document) -> None:
    """Draw footer only on page 1."""
    W, H = A4
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#F7FAFC"))
    canvas.rect(0, 0, W, 0.52 * inch, stroke=0, fill=1)
    canvas.setStrokeColor(colors.HexColor("#CBD5E0"))
    canvas.setLineWidth(0.5)
    canvas.line(0.68 * inch, 0.52 * inch, W - 0.68 * inch, 0.52 * inch)
    canvas.setFillColor(colors.HexColor(f"#{_SLATE}"))
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(0.68 * inch, 0.28 * inch, f"Generated by Agent OTG  •  DWE Team  •  {_TODAY}")
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.drawRightString(W - 0.68 * inch, 0.28 * inch, f"Page {document.page}")
    canvas.restoreState()


def _pdf_header_footer(canvas, document, title: str) -> None:
    """Draw header + footer on subsequent pages."""
    W, H = A4
    canvas.saveState()

    # Header bar (thin teal stripe)
    canvas.setFillColor(colors.HexColor(f"#{_TEAL}"))
    canvas.rect(0, H - 0.32 * inch, W, 0.32 * inch, stroke=0, fill=1)

    # Header text
    canvas.setFillColor(colors.white)
    if LOGO_PATH.is_file():
        canvas.drawImage(str(LOGO_PATH), 0.40 * inch, H - 0.285 * inch,
                         width=0.19 * inch, height=0.19 * inch, mask="auto", preserveAspectRatio=True)
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.drawString(0.68 * inch, H - 0.22 * inch, "★ AGENT OTG  |  DWE TEAM")
    canvas.setFont("Helvetica", 7.5)
    short_title = (title[:52] + "…") if len(title) > 55 else title
    canvas.drawRightString(W - 0.68 * inch, H - 0.22 * inch, short_title)

    # Footer background
    canvas.setFillColor(colors.HexColor("#F7FAFC"))
    canvas.rect(0, 0, W, 0.52 * inch, stroke=0, fill=1)

    # Footer top border line
    canvas.setStrokeColor(colors.HexColor("#CBD5E0"))
    canvas.setLineWidth(0.5)
    canvas.line(0.68 * inch, 0.52 * inch, W - 0.68 * inch, 0.52 * inch)

    # Footer text
    canvas.setFillColor(colors.HexColor(f"#{_SLATE}"))
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(0.68 * inch, 0.28 * inch, f"Generated by Agent OTG  •  {_TODAY}")
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.drawRightString(W - 0.68 * inch, 0.28 * inch, f"Page {document.page}")

    canvas.restoreState()


def _parse_markdown_table(lines: list[str]) -> list[list[str]] | None:
    """Parse a markdown table from a list of lines. Returns rows or None."""
    if len(lines) < 2:
        return None
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            return None
        # Skip separator rows like |---|---|
        if re.match(r"^\|[-| :]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells:
            rows.append(cells)
    return rows if len(rows) >= 1 else None


def _build_pdf_table(rows: list[list[str]]) -> Table:
    """Build a styled ReportLab Table from parsed markdown rows."""
    col_count = max(len(r) for r in rows)
    # Normalize row widths
    normalized = [r + [""] * (col_count - len(r)) for r in rows]

    W = A4[0] - 1.36 * inch  # available width

    col_width = W / max(col_count, 1)
    col_widths = [col_width] * col_count

    table = Table(normalized, colWidths=col_widths, repeatRows=1)
    style = TableStyle([
        # Header row
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_TEAL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
        ("TOPPADDING", (0, 0), (-1, 0), 7),
        # Body rows
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor(f"#{_INK}")),
        # Alternating row shading
        *[("BACKGROUND", (0, i), (-1, i), colors.HexColor(f"#{_TEAL_LIGHT}"))
          for i in range(2, len(normalized), 2)],
        # Grid
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E0")),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor(f"#{_TEAL_DARK}")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{_TEAL_LIGHT}")]),
        # Alignment
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])
    table.setStyle(style)
    return table


def _generate_pdf(title: str, content: str, path: Path) -> None:
    """Advanced PDF generation with cover header, tables, code blocks, callouts."""
    styles = _make_styles()

    def _first_page(canvas, document):
        _pdf_cover_page(canvas, document, title)
        _pdf_footer_only(canvas, document)

    def _later_pages(canvas, document):
        _pdf_header_footer(canvas, document, title)

    # Build story
    story: list = []

    # ── Initial spacer to clear executive top band cleanly ───────────────────
    # Cover band is 2.45", topMargin is 0.72", so 1.85" spacer positions content right below it
    story.append(Spacer(1, 1.85 * inch))

    # ── Parse and render content lines ───────────────────────────────────────
    lines = (content or "").splitlines()
    i = 0
    _numbered_counter = [0]

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # ── Fenced code block ─────────────────────────────────────────────
        if line.strip().startswith("```"):
            code_lines: list[str] = []
            lang = line.strip()[3:].strip()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            code_text = "\n".join(code_lines) or " "
            label = f"[{lang.upper()}]" if lang else "[CODE]"
            story.append(Spacer(1, 4))
            story.append(Paragraph(f"<font color='#{_TEAL}' size='8'>{_escape(label)}</font>", styles["body"]))
            story.append(Preformatted(code_text, styles["code"]))
            i += 1
            continue

        # ── Markdown table detection ──────────────────────────────────────
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parsed = _parse_markdown_table(table_lines)
            if parsed:
                story.append(Spacer(1, 6))
                story.append(_build_pdf_table(parsed))
                story.append(Spacer(1, 8))
            else:
                for tl in table_lines:
                    story.append(Paragraph(_escape(tl.strip("|").strip()), styles["body"]))
            continue

        # ── Callouts (> Note: ... / > Warning: ...) ───────────────────────
        if line.startswith("> "):
            inner = line[2:].strip()
            if inner.lower().startswith("warning:") or inner.lower().startswith("caution:"):
                story.append(Paragraph("⚠  " + _escape(inner), styles["callout_warn"]))
            else:
                story.append(Paragraph("ℹ  " + _escape(inner), styles["callout_note"]))
            i += 1
            continue

        # ── Headings ──────────────────────────────────────────────────────
        if line.startswith("### "):
            _numbered_counter[0] = 0
            story.append(Paragraph(_escape(line[4:].strip()), styles["h3"]))
        elif line.startswith("## "):
            _numbered_counter[0] = 0
            story.append(Paragraph(_escape(line[3:].strip()), styles["h2"]))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E0"), spaceAfter=4))
        elif line.startswith("# "):
            _numbered_counter[0] = 0
            story.append(Paragraph(_escape(line[2:].strip()), styles["h1"]))
            story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor(f"#{_TEAL}"), spaceAfter=6))

        # ── Horizontal rule ───────────────────────────────────────────────
        elif line.strip() in ("---", "***", "___"):
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E0"), spaceBefore=6, spaceAfter=6))

        # ── Bullet list (- or *) ──────────────────────────────────────────
        elif re.match(r"^(\s{2,4})?[-*]\s", line):
            indent = len(line) - len(line.lstrip())
            text = re.sub(r"^\s*[-*]\s", "", line).strip()
            # Bold/italic inline
            text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _escape(text))
            text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
            sty = styles["bullet2"] if indent >= 2 else styles["bullet"]
            story.append(Paragraph(f"<font color='#{_TEAL}'>•</font>  {text}", sty))

        # ── Numbered list ─────────────────────────────────────────────────
        elif re.match(r"^\d+\.\s", line):
            _numbered_counter[0] += 1
            text = re.sub(r"^\d+\.\s+", "", line).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _escape(text))
            text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
            story.append(Paragraph(
                f"<font color='#{_TEAL}' size='10'><b>{_numbered_counter[0]}.</b></font>  {text}",
                styles["numbered"],
            ))

        # ── Empty line → small spacer ──────────────────────────────────────
        elif not line.strip():
            story.append(Spacer(1, 4))

        # ── Regular body paragraph ─────────────────────────────────────────
        else:
            _numbered_counter[0] = 0
            text = _escape(line.strip())
            # Inline bold/italic
            text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
            text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
            text = re.sub(r"`(.+?)`", r"<font name='Courier' size='9' color='#2D3748'>\1</font>", text)
            story.append(Paragraph(text, styles["body"]))

        i += 1

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        title=title,
        author="Agent OTG | DWE Team",
        subject="Generated by Agent OTG",
        leftMargin=0.68 * inch,
        rightMargin=0.68 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.72 * inch,
    )
    doc.build(story, onFirstPage=_first_page, onLaterPages=_later_pages)


# ══════════════════════════════════════════════════════════════════════════════
# DOCX GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def _add_paragraph_run(p, text: str, bold=False, italic=False, is_code=False,
                        font_name="Calibri", size_pt=11,
                        color_hex: str | None = None) -> None:
    """Add a formatted run to a python-docx paragraph with code, bold, and italic support."""
    run = p.add_run(text)
    run.bold  = bold
    run.italic = italic
    if is_code:
        run.font.name = "Consolas"
        run.font.size = Pt(size_pt - 0.5)
        run.font.color.rgb = RGBColor.from_string("1E293B")
    else:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        if color_hex:
            run.font.color.rgb = RGBColor.from_string(color_hex.lstrip("#"))


def _add_shading_to_paragraph(p, fill_hex: str) -> None:
    """Give a paragraph a solid background color (for code blocks)."""
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    pPr.append(shd)


def _add_cell_shading(cell, fill_hex: str) -> None:
    """Give a table cell a solid background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    tcPr.append(shd)


def _build_docx_table(doc: Document, rows: list[list[str]], total_width_in: float = 6.5) -> None:
    """Render a styled table in python-docx with brand colors, alternating fills, and distributed column widths."""
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    if col_count == 0:
        return
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.autofit = False

    col_width = Inches(total_width_in / col_count)
    for col in table.columns:
        col.width = col_width

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_width
            val = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(4)
            p.paragraph_format.line_spacing = Pt(13)
            if r_idx == 0:
                _add_cell_shading(cell, _TEAL_DARK)
                _add_paragraph_run(p, val.strip(), bold=True, color_hex="FFFFFF", font_name="Calibri", size_pt=10)
            else:
                if r_idx % 2 == 1:
                    _add_cell_shading(cell, _TEAL_LIGHT)
                _add_paragraph_run(p, val.strip(), font_name="Calibri", size_pt=9.5, color_hex=_INK)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after  = Pt(4)


def _add_left_border(p, color_hex: str, width_pt: int = 12) -> None:
    """Add a colored left-border accent line to a paragraph (callout style)."""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(width_pt))
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(left)
    pPr.append(pBdr)


def _inline_parse(text: str) -> list[tuple[str, bool, bool, bool]]:
    """
    Parse **bold**, *italic*, and `code` inline markdown and return a list of
    (text_fragment, is_bold, is_italic, is_code) tuples.
    """
    parts: list[tuple[str, bool, bool, bool]] = []
    remaining = text
    while remaining:
        bold_m   = re.search(r"\*\*(.+?)\*\*", remaining)
        italic_m = re.search(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", remaining)
        code_m   = re.search(r"`([^`]+)`", remaining)

        matches = [m for m in [bold_m, italic_m, code_m] if m]
        if not matches:
            if remaining:
                parts.append((remaining, False, False, False))
            break

        first = min(matches, key=lambda m: m.start())
        if first.start() > 0:
            parts.append((remaining[:first.start()], False, False, False))

        inner = first.group(1)
        is_bold   = (first == bold_m)
        is_italic = (first == italic_m)
        is_code   = (first == code_m)
        parts.append((inner, is_bold, is_italic, is_code))
        remaining = remaining[first.end():]

    return parts or [(text, False, False, False)]


def _generate_docx(title: str, content: str, path: Path) -> None:
    """Professional Word document with clean structure, branding without logo, and typography."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)

    # ── Base styles ──────────────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(_INK)
    normal.paragraph_format.space_after  = Pt(6)
    normal.paragraph_format.line_spacing = Pt(14.5)

    # Branded masthead bar (NO logo image, branding text only as requested)
    brand_table = doc.add_table(rows=1, cols=1)
    brand_table.autofit = False
    brand_table.columns[0].width = Inches(6.5)
    b_cell = brand_table.cell(0, 0)
    b_cell.width = Inches(6.5)
    _add_cell_shading(b_cell, _TEAL_DARK)
    brand_p = b_cell.paragraphs[0]
    brand_p.paragraph_format.space_before = Pt(6)
    brand_p.paragraph_format.space_after  = Pt(6)
    brand_p.paragraph_format.left_indent  = Inches(0.12)
    _add_paragraph_run(brand_p, "AGENT OTG  |  DWE TEAM", bold=True, font_name="Calibri", size_pt=10, color_hex=_WHITE)

    # ── Document Title ───────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after  = Pt(4)
    title_run = title_p.add_run(title)
    title_run.font.name  = "Calibri"
    title_run.font.size  = Pt(26)
    title_run.font.bold  = True
    title_run.font.color.rgb = RGBColor.from_string(_TEAL_DARK)

    # Accent underline rule under title
    rule_p = doc.add_paragraph()
    rule_p.paragraph_format.space_before = Pt(0)
    rule_p.paragraph_format.space_after  = Pt(8)
    pPr = rule_p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), _TEAL)
    pBdr.append(bot)
    pPr.append(pBdr)

    # Subtitle line
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    _add_paragraph_run(sub_p, f"Agent OTG  •  DWE Team  •  {_TODAY}",
                       font_name="Calibri", size_pt=9.5, color_hex=_SLATE)

    # ── Content: line-by-line markdown renderer ───────────────────────────────
    lines = (content or "").splitlines()
    i = 0
    in_code_block = False
    code_lines: list[str] = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # ── Code block ────────────────────────────────────────────────────
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Code block container with left teal border
                code_tbl = doc.add_table(rows=1, cols=1)
                code_tbl.autofit = False
                code_tbl.columns[0].width = Inches(6.5)
                c_cell = code_tbl.cell(0, 0)
                c_cell.width = Inches(6.5)
                _add_cell_shading(c_cell, "F4F6F9")
                _add_left_border(c_cell.paragraphs[0], _TEAL, width_pt=18)
                for c_idx, cl in enumerate(code_lines):
                    cp = c_cell.paragraphs[0] if c_idx == 0 else c_cell.add_paragraph()
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after  = Pt(1.5)
                    cp.paragraph_format.left_indent  = Inches(0.12)
                    _add_paragraph_run(cp, cl or " ", font_name="Consolas", size_pt=9, color_hex="1E293B")
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after  = Pt(4)
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ── Markdown table ────────────────────────────────────────────────
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parsed = _parse_markdown_table(table_lines)
            if parsed:
                _build_docx_table(doc, parsed, total_width_in=6.5)
            continue

        stripped_line = line.strip()

        # ── Skip duplicate title or redundant top heading if repeated at top ──
        if (i == 0 or (i < 3 and not stripped_line.startswith(("-", "*", ">", "|", "```")))) and stripped_line.startswith("# "):
            h_text = stripped_line[2:].strip().lower()
            if h_text == title.strip().lower() or h_text in ("introduction", "overview", "summary", "document", "report"):
                i += 1
                continue

        # ── Headings ─────────────────────────────────────────────────────
        if line.startswith("### "):
            text = re.sub(r"[\*`_]", "", line[4:]).strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(10)
            hp.paragraph_format.space_after  = Pt(2)
            hp.paragraph_format.keep_with_next = True
            _add_paragraph_run(hp, text or " ", bold=True, font_name="Calibri", size_pt=11.5, color_hex=_INK)

        elif line.startswith("## "):
            text = re.sub(r"[\*`_]", "", line[3:]).strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(14)
            hp.paragraph_format.space_after  = Pt(3)
            hp.paragraph_format.keep_with_next = True
            _add_paragraph_run(hp, text or " ", bold=True, font_name="Calibri", size_pt=13.5, color_hex=_TEAL)

        elif line.startswith("# "):
            text = re.sub(r"[\*`_]", "", line[2:]).strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(18)
            hp.paragraph_format.space_after  = Pt(4)
            hp.paragraph_format.keep_with_next = True
            pPr_h = hp._p.get_or_add_pPr()
            pBdr_h = OxmlElement("w:pBdr")
            bot_h = OxmlElement("w:bottom")
            bot_h.set(qn("w:val"), "single")
            bot_h.set(qn("w:sz"), "6")
            bot_h.set(qn("w:space"), "2")
            bot_h.set(qn("w:color"), "CBD5E0")
            pBdr_h.append(bot_h)
            pPr_h.append(pBdr_h)
            _add_paragraph_run(hp, text or " ", bold=True, font_name="Calibri", size_pt=16, color_hex=_TEAL_DARK)

        # ── Horizontal rule ───────────────────────────────────────────────
        elif line.strip() in ("---", "***", "___"):
            hr_p = doc.add_paragraph()
            hr_p.paragraph_format.space_before = Pt(4)
            hr_p.paragraph_format.space_after  = Pt(4)
            pPr2 = hr_p._p.get_or_add_pPr()
            pBdr2 = OxmlElement("w:pBdr")
            bot2 = OxmlElement("w:bottom")
            bot2.set(qn("w:val"), "single")
            bot2.set(qn("w:sz"), "6")
            bot2.set(qn("w:space"), "1")
            bot2.set(qn("w:color"), "CBD5E0")
            pBdr2.append(bot2)
            pPr2.append(pBdr2)

        # ── Callouts  > Note / > Warning ─────────────────────────────────
        elif line.startswith("> "):
            inner = line[2:].strip()
            is_warn = inner.lower().startswith(("warning:", "caution:"))
            color_hex = _WARN if is_warn else _NOTE
            bg_hex    = "FFF9F0" if is_warn else "EBF8FF"
            icon      = "⚠  " if is_warn else "ℹ  "
            cp = doc.add_paragraph()
            cp.paragraph_format.left_indent  = Inches(0.2)
            cp.paragraph_format.space_before = Pt(4)
            cp.paragraph_format.space_after  = Pt(6)
            _add_shading_to_paragraph(cp, bg_hex)
            _add_left_border(cp, color_hex, width_pt=18)
            for frag, b, it, code in _inline_parse(icon + inner):
                _add_paragraph_run(cp, frag, bold=b, italic=it, is_code=code,
                                   font_name="Calibri", size_pt=10.5, color_hex=color_hex)

        # ── Bullet list ───────────────────────────────────────────────────
        elif re.match(r"^\s{0,4}[-*]\s", line):
            indent = len(line) - len(line.lstrip())
            text   = re.sub(r"^\s*[-*]\s+", "", line).strip()
            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent  = Inches(0.4 if indent >= 2 else 0.25)
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after  = Pt(2.5)
            bp.paragraph_format.line_spacing = Pt(14)
            _add_paragraph_run(bp, "•  ", bold=True, font_name="Calibri", size_pt=10.5, color_hex=_TEAL)
            for frag, b, it, code in _inline_parse(text):
                _add_paragraph_run(bp, frag, bold=b, italic=it, is_code=code, font_name="Calibri", size_pt=10.5, color_hex=_INK)

        # ── Numbered list ─────────────────────────────────────────────────
        elif re.match(r"^\s*\d+[\.\)]\s", line):
            m = re.match(r"^\s*(\d+)[\.\)]\s+(.*)$", line)
            num = m.group(1) if m else "1"
            text = m.group(2).strip() if m else line.strip()
            np_ = doc.add_paragraph()
            np_.paragraph_format.left_indent  = Inches(0.25)
            np_.paragraph_format.space_before = Pt(1)
            np_.paragraph_format.space_after  = Pt(2.5)
            np_.paragraph_format.line_spacing = Pt(14)
            _add_paragraph_run(np_, f"{num}.  ", bold=True, font_name="Calibri", size_pt=10.5, color_hex=_TEAL)
            for frag, b, it, code in _inline_parse(text):
                _add_paragraph_run(np_, frag, bold=b, italic=it, is_code=code, font_name="Calibri", size_pt=10.5, color_hex=_INK)

        # ── Empty line (spacing already handled by paragraph space_after) ─
        elif not line.strip():
            pass

        # ── Body paragraph ────────────────────────────────────────────────
        else:
            bp = doc.add_paragraph()
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after  = Pt(6)
            bp.paragraph_format.line_spacing = Pt(14.5)
            for frag, b, it, code in _inline_parse(line.strip()):
                _add_paragraph_run(bp, frag, bold=b, italic=it, is_code=code, font_name="Calibri", size_pt=10.5, color_hex=_INK)

        i += 1

    # Flush unclosed code block if content ended inside code fence
    if in_code_block and code_lines:
        code_tbl = doc.add_table(rows=1, cols=1)
        code_tbl.autofit = False
        code_tbl.columns[0].width = Inches(6.5)
        c_cell = code_tbl.cell(0, 0)
        c_cell.width = Inches(6.5)
        _add_cell_shading(c_cell, "F4F6F9")
        _add_left_border(c_cell.paragraphs[0], _TEAL, width_pt=18)
        for c_idx, cl in enumerate(code_lines):
            cp = c_cell.paragraphs[0] if c_idx == 0 else c_cell.add_paragraph()
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after  = Pt(1.5)
            cp.paragraph_format.left_indent  = Inches(0.12)
            _add_paragraph_run(cp, cl or " ", font_name="Consolas", size_pt=9, color_hex="1E293B")

    # ── Footer ────────────────────────────────────────────────────────────────
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer.add_run(f"Agent OTG  |  DWE Team  |  {_TODAY}")
    f_run.font.size = Pt(8)
    f_run.font.color.rgb = RGBColor.from_string(_SLATE)
    f_run.font.name = "Calibri"

    doc.save(path)


# ══════════════════════════════════════════════════════════════════════════════
# PPTX GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def _generate_pptx(title: str, content: str, path: Path) -> None:
    """Professional PPTX with title slide + content slides using consistent teal brand."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor as PR
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches as I, Pt as P, Emu
    except ImportError as exc:
        raise RuntimeError("PPTX requires python-pptx: pip install python-pptx") from exc

    prs = Presentation()
    prs.slide_width  = I(13.333)
    prs.slide_height = I(7.5)

    TEAL  = PR(8,  126, 112)
    DARK  = PR(6,  95,  84)
    INK   = PR(24, 33,  43)
    LIGHT = PR(232,245, 243)
    SLATE = PR(74, 85,  104)
    WHITE = PR(255,255, 255)
    GRAY  = PR(248,250, 249)

    def _blank_slide() -> object:
        return prs.slides.add_slide(prs.slide_layouts[6])

    def _set_bg(slide, color: PR):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = color

    def _add_shape(slide, l, t, w, h, color: PR, line=False):
        from pptx.util import Emu
        shape = slide.shapes.add_shape(1, I(l), I(t), I(w), I(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if not line:
            shape.line.fill.background()
        return shape

    def _add_text(slide, text, l, t, w, h,
                  font_name="Calibri", size=18, bold=False,
                  color=None, align=PP_ALIGN.LEFT, wrap=True):
        tb = slide.shapes.add_textbox(I(l), I(t), I(w), I(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        p  = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.name = font_name
        run.font.size = P(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        return tf

    def _add_footer(slide, page_num: int):
        _add_text(slide, f"Agent OTG  |  DWE Team  |  {_TODAY}  |  {page_num}",
                  0.4, 7.15, 12.5, 0.25,
                  font_name="Calibri", size=8, color=SLATE, align=PP_ALIGN.CENTER)

    def _add_logo(slide, left: float, top: float, size: float = 0.30):
        if LOGO_PATH.is_file():
            slide.shapes.add_picture(str(LOGO_PATH), I(left), I(top), width=I(size), height=I(size))

    # ── SLIDE 1: Title slide ────────────────────────────────────────────────
    s0 = _blank_slide()
    _set_bg(s0, GRAY)

    # Left teal panel
    _add_shape(s0, 0, 0, 4.8, 7.5, DARK)
    # Accent stripe
    _add_shape(s0, 4.8, 0, 0.07, 7.5, TEAL)

    # Title text on left panel
    _add_text(s0, "AGENT OTG", 0.35, 0.4, 4.1, 0.45,
              font_name="Calibri", size=11, bold=True, color=PR(160, 212, 206))
    _add_logo(s0, 3.85, 0.25, 0.58)
    _add_text(s0, title, 0.35, 1.0, 4.1, 3.2,
              font_name="Calibri", size=30, bold=True, color=WHITE)
    _add_text(s0, f"DWE Team  •  {_TODAY}", 0.35, 6.6, 4.1, 0.4,
              font_name="Calibri", size=10, color=PR(160, 212, 206))

    # Right panel — executive deck card
    _add_shape(s0, 5.5, 1.3, 7.0, 4.9, WHITE)
    _add_shape(s0, 5.5, 1.3, 0.09, 4.9, TEAL)
    _add_text(s0, "EXECUTIVE PRESENTATION BRIEF", 5.9, 1.6, 6.2, 0.4,
              font_name="Calibri", size=12, bold=True, color=DARK)
    _add_text(s0, f"• Platform: Agent OTG — Team DWE\n"
                  f"• Generated: {_TODAY}\n"
                  f"• Architecture: Local Multi-Model Intelligence\n"
                  f"• Display: 16:9 Widescreen Executive Layout\n"
                  f"• Status: Offline Verified",
              5.9, 2.2, 6.2, 3.4,
              font_name="Calibri", size=15, bold=False, color=INK)
    _add_footer(s0, 1)

    # ── Parse content into slides ───────────────────────────────────────────
    # Split on markdown headings or common model forms such as
    # "Slide 2: Deployment".  Without this normalization a valid-looking
    # model response can become one crowded slide containing the entire deck.
    raw_lines = (content or "").splitlines()
    slides_data: list[tuple[str, list[str]]] = []
    current_title = f"{title} — Executive Summary"
    current_body: list[str] = []

    for line in raw_lines:
        stripped = line.strip()
        slide_label = re.match(r"^(?:slide\s*\d+\s*[:\-–—]\s*|slide\s*\d+\s+)(.+)$", stripped, re.I)
        if stripped.startswith("# ") or stripped.startswith("## ") or slide_label:
            if current_body or slides_data:
                slides_data.append((current_title, current_body))
            h_raw = slide_label.group(1).strip() if slide_label else re.sub(r"^#+\s+", "", stripped)
            h_raw = re.sub(r"[\*`_]", "", h_raw).strip()
            if h_raw.lower() in ("introduction", "intro", "overview"):
                h_raw = f"{title} — Executive Summary"
            elif h_raw.lower() == title.lower():
                h_raw = f"{title} — Overview"
            current_title = h_raw
            current_body  = []
        else:
            if stripped:  # skip blank lines at top of slide
                current_body.append(stripped)

    if current_title or current_body:
        slides_data.append((current_title, current_body))

    if not slides_data:
        slides_data = [(title, ["No content was generated."])]

    # If a model returned only a flat bullet list, create a readable deck
    # rather than silently rendering one giant "Overview" slide.
    if len(slides_data) == 1 and slides_data[0][0] in ("Overview", f"{title} — Executive Summary") and len(slides_data[0][1]) > 6:
        flat_lines = slides_data[0][1]
        slides_data = [
            (title if number == 0 else f"{title} — Key Points {number + 1}", flat_lines[start:start + 5])
            for number, start in enumerate(range(0, len(flat_lines), 5))
        ]

    # ── Generate content slides ─────────────────────────────────────────────
    # A fixed-size slide cannot remain readable with 18 long bullets.  Split
    # oversized sections into continuation slides so text stays on-canvas.
    paged_slides: list[tuple[str, list[str]]] = []
    for slide_title, body_lines in slides_data:
        lines = body_lines or ["No content was generated."]
        for page, start in enumerate(range(0, len(lines), 6), 1):
            suffix = "" if page == 1 else f" (continued {page})"
            paged_slides.append((f"{slide_title}{suffix}", lines[start:start + 6]))

    fallback_slide_titles = [
        "Executive Summary", "Context and Objectives", "Core Concepts", "Evidence and Examples",
        "Strategic Implications", "Risks and Constraints", "Implementation Roadmap", "Recommendations",
    ]
    for slide_num, (slide_title, body_lines) in enumerate(paged_slides, 2):
        # Protect direct tool calls too: a raw `Slide_1` is never a
        # professional presentation heading.
        if re.fullmatch(r"(?:slide[_\s-]*\d+|slide\s*title|overview)", slide_title.strip(), re.I):
            slide_title = f"{title} — {fallback_slide_titles[(slide_num - 2) % len(fallback_slide_titles)]}"
        sl = _blank_slide()
        _set_bg(sl, GRAY)

        # Top teal header band
        _add_shape(sl, 0, 0, 13.333, 1.2, DARK)
        # Accent stripe at bottom of band
        _add_shape(sl, 0, 1.2, 13.333, 0.05, TEAL)

        # Slide number badge (top left)
        _add_text(sl, str(slide_num - 1), 0.22, 0.35, 0.5, 0.5,
                  font_name="Calibri", size=9, bold=True, color=PR(160, 212, 206))

        # Slide title in header
        _add_text(sl, slide_title, 0.75, 0.15, 11.8, 0.9,
                  font_name="Calibri", size=24, bold=True, color=WHITE)

        # "Agent OTG" label (top right)
        _add_text(sl, "AGENT OTG", 11.2, 0.0, 2.0, 0.35,
                  font_name="Calibri", size=8, bold=True,
                  color=PR(160, 212, 206), align=PP_ALIGN.RIGHT)
        _add_logo(sl, 12.72, 0.39, 0.22)

        # Body content area
        content_tf = slide.shapes.add_textbox(
            I(0.55), I(1.38), I(12.2), I(5.6)
        ).text_frame if False else None  # unused

        body_tb = sl.shapes.add_textbox(I(0.55), I(1.38), I(12.2), I(5.6))
        body_tf = body_tb.text_frame
        body_tf.word_wrap = True

        first_para = True
        for bline in body_lines:
            bline = bline.strip()
            if not bline:
                continue
            # Clean leading markers
            is_bullet = bline.startswith(("- ", "* ", "• "))
            is_numbered = bool(re.match(r"^\d+\.\s", bline))
            text = re.sub(r"^[-*•]\s+", "", bline)
            text = re.sub(r"^\d+\.\s+", "", text)
            # Strip markdown bold for PPT (apply bold via font)
            is_bold = bool(re.match(r"^\*\*.*\*\*$", text))
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*",    r"\1", text)
            text = re.sub(r"`(.+?)`",      r"\1", text)

            if first_para:
                para = body_tf.paragraphs[0]
                first_para = False
            else:
                para = body_tf.add_paragraph()

            if is_bullet or is_numbered:
                para.level = 0
                bullet_char = "▸ " if is_bullet else ""
                run = para.add_run()
                run.text = bullet_char + text
                run.font.name = "Calibri"
                run.font.size = P(16)
                run.font.bold = is_bold
                run.font.color.rgb = INK
                para.space_after = P(6)
            else:
                run = para.add_run()
                run.text = text
                run.font.name = "Calibri"
                run.font.size = P(16)
                run.font.bold = is_bold or bline.startswith("###")
                run.font.color.rgb = DARK if is_bold else INK
                para.space_after = P(4)

        # Left accent bar
        _add_shape(sl, 0, 1.28, 0.06, 5.9, TEAL)

        _add_footer(sl, slide_num)

    prs.save(path)


# ══════════════════════════════════════════════════════════════════════════════
# XLSX GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def _generate_xlsx(title: str, content: str, path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "", title)[:31] or "Sheet1"
    rows = _rows(content)

    thin = Side(style="thin", color="CBD5E0")
    thick = Side(style="medium", color=_TEAL)

    # 1. Merged Title Bar on Row 1
    col_count = max(1, len(rows[0]))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    title_cell = sheet.cell(1, 1, title[:110])
    title_cell.font = Font(name="Calibri", bold=True, color=_WHITE, size=13)
    title_cell.fill = PatternFill("solid", fgColor=_TEAL_DARK)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    if LOGO_PATH.is_file():
        from openpyxl.drawing.image import Image as ExcelImage
        logo = ExcelImage(str(LOGO_PATH))
        logo.width = 30
        logo.height = 30
        sheet.add_image(logo, f"{get_column_letter(col_count)}1")

    # 2. Header Row on Row 2
    header_row = rows[0]
    sheet.append(header_row)
    for col_idx, cell in enumerate(sheet[2], 1):
        cell.font = Font(name="Calibri", bold=True, color=_WHITE, size=11)
        cell.fill = PatternFill("solid", fgColor=_TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thick, left=thin, right=thin, top=thin)
    sheet.row_dimensions[2].height = 24

    # 3. Data Rows on Row 3+ with native int/float types
    for row_idx, row_data in enumerate(rows[1:], start=3):
        parsed_row = [_parse_cell_value(c) for c in row_data]
        sheet.append(parsed_row)
        bg = _TEAL_LIGHT if row_idx % 2 == 0 else _WHITE
        for col_idx, cell in enumerate(sheet[row_idx], 1):
            cell.font = Font(name="Calibri", size=10.5, color=_INK)
            cell.fill = PatternFill("solid", fgColor=bg)
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(cell.value, bool):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin, left=thin, right=thin, top=thin)
        sheet.row_dimensions[row_idx].height = 20

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    # 4. Auto-fit column widths using header & data rows ONLY (rows 2 to max_row)
    for column in range(1, sheet.max_column + 1):
        col_letter = get_column_letter(column)
        max_len = 0
        for r in range(2, sheet.max_row + 1):
            val_str = str(sheet.cell(r, column).value or "")
            max_len = max(max_len, len(val_str))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    workbook.save(path)
    if load_workbook(path, read_only=True).active.max_row < 1:
        raise RuntimeError("Generated XLSX has no rows")


def _media_type(extension: str) -> str:
    return {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "csv": "text/csv",
        "json": "application/json",
        "md": "text/markdown",
        "txt": "text/plain",
    }[extension]


FORMAT_ALIASES = {
    "doc": "docx", "word": "docx", "msword": "docx", "worddoc": "docx",
    "excel": "xlsx", "spreadsheet": "xlsx", "sheet": "xlsx",
    "ppt": "pptx", "powerpoint": "pptx", "slides": "pptx",
    "text": "txt", "markdown": "md",
}

def create_artifact(title: str, content: str, file_format: str, filename: str | None = None) -> Artifact:
    extension = str(file_format).lower().lstrip(".")
    extension = FORMAT_ALIASES.get(extension, extension)
    if extension not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported format '{file_format}'. Supported: {', '.join(sorted(SUPPORTED_FORMATS))}")

    # Always derive clean, AI-decided document title and filename slug
    clean_t, clean_c, safe_slug = derive_clean_title(title, content)
    title = clean_t
    content = clean_c

    # File name MUST be derived strictly from the AI-decided title
    if not filename or any(p in filename.lower() for p in ["output", "document", "give", "write", "create", "generate", "/agent", "prompt", "test_alias"]):
        filename = f"{safe_slug}.{extension}"
    else:
        # Check if caller passed a raw prompt or generic name as filename
        f_stem = Path(filename).stem
        if any(p in f_stem.lower() for p in ["give", "write", "create", "generate", "agent", "prompt", "output", "document"]):
            filename = f"{safe_slug}.{extension}"
        else:
            filename = f"{re.sub(r'[^a-zA-Z0-9._-]+', '_', f_stem).strip('_')}.{extension}"

    path = _safe_filename(filename, extension)
    if extension == "pdf":
        _generate_pdf(title, content, path)
        if len(PdfReader(str(path)).pages) < 1:
            raise RuntimeError("Generated PDF has no pages")
    elif extension == "docx":
        _generate_docx(title, content, path)
        if not Document(path).paragraphs:
            raise RuntimeError("Generated DOCX has no paragraphs")
    elif extension == "pptx":
        _generate_pptx(title, content, path)
    elif extension == "xlsx":
        _generate_xlsx(title, content, path)
    elif extension == "csv":
        with path.open("w", encoding="utf-8", newline="") as stream:
            stream.write(content)
    elif extension == "json":
        try:
            payload = json.loads(content)
        except json.JSONDecodeError:
            payload = {"title": title, "content": content}
        path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        path.write_text(content, encoding="utf-8")

    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Artifact validation failed for {path.name}")
    return Artifact(path=path, media_type=_media_type(extension), size_bytes=path.stat().st_size)



def resolve_artifact(filename: str) -> Path:
    path = (OUTPUT_DIR / Path(filename).name).resolve()
    if path.parent != OUTPUT_DIR or not path.is_file():
        raise FileNotFoundError(filename)
    return path
