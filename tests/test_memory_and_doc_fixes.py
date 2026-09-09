import os
from pathlib import Path
from unittest.mock import MagicMock
import openpyxl
from docx import Document

import router
import langgraph_agent
from artifacts import OUTPUT_DIR, create_artifact, _rows, _parse_cell_value, _generate_xlsx
import tools


def test_router_history_retains_at_least_10_turns():
    """Verify that router.add_to_history does not prematurely purge messages after 3 turns."""
    router.clear_history()
    # Add 8 full turns (16 messages)
    for i in range(1, 9):
        router.add_to_history("user", f"Question {i}")
        router.add_to_history("assistant", f"Answer {i}")

    history = router.get_history()
    assert len(history) == 16
    assert history[0]["content"] == "Question 1"
    assert history[-1]["content"] == "Answer 8"

    # Add up to 12 turns (24 messages). Max messages is max(MAX_TURNS * 2, 20) = 20 (if MAX_TURNS=6)
    for i in range(9, 13):
        router.add_to_history("user", f"Question {i}")
        router.add_to_history("assistant", f"Answer {i}")

    assert len(router.get_history()) >= 20


def test_clear_history_creates_new_session():
    """Verify clear_history clears memory and updates session name."""
    router.add_to_history("user", "Hello")
    old_session = router._CURRENT_SESSION_NAME
    router.clear_history()
    assert len(router.get_history()) == 0
    assert router._CURRENT_SESSION_NAME != ""


def test_referential_topic_resolution_from_history():
    """Verify that 'put it in a word file' resolves subject from prior turn."""
    history = [
        {"role": "user", "content": "Tell me about Quantum Computing and its future"},
        {"role": "assistant", "content": "✅ **quantum_computing_overview.pdf** created successfully!\n📁 Saved to disk."},
    ]
    resolved, source = langgraph_agent._resolve_subject_from_history("Now put it in a word file", history, "docx")
    assert "Quantum Computing" in resolved


def test_referential_topic_resolution_from_user_turn():
    """Verify that 'save this as excel' resolves subject from previous user question."""
    history = [
        {"role": "user", "content": "Analyze renewable solar energy output for 2026"},
        {"role": "assistant", "content": "Here is the analysis on solar energy: panels, efficiency, grid tie-in."},
    ]
    resolved, source = langgraph_agent._resolve_subject_from_history("Save this as excel", history, "xlsx")
    assert "Solar Energy" in resolved or "Renewable Solar" in resolved


def test_excel_rows_parses_markdown_table():
    """Verify that markdown tables are parsed into real multi-column rows."""
    md_table = (
        "| Employee ID | Name | Department | Salary |\n"
        "|-------------|------|------------|--------|\n"
        "| EMP-001     | Alice | Engineering| 95000  |\n"
        "| EMP-002     | Bob   | Marketing  | 78000  |\n"
    )
    rows = _rows(md_table)
    assert len(rows) == 3  # header + 2 data rows
    assert rows[0] == ["Employee ID", "Name", "Department", "Salary"]
    assert rows[1] == ["EMP-001", "Alice", "Engineering", "95000"]
    assert rows[2] == ["EMP-002", "Bob", "Marketing", "78000"]


def test_excel_rows_strips_code_fences_and_preamble():
    """Verify fenced code blocks and conversational preamble are stripped."""
    fenced_content = (
        "Here is the spreadsheet you asked for:\n"
        "```csv\n"
        "Product,Units Sold,Revenue\n"
        "Widget A,150,4500\n"
        "Widget B,80,2400\n"
        "```\n"
        "Hope this helps!"
    )
    rows = _rows(fenced_content)
    assert len(rows) == 3
    assert rows[0] == ["Product", "Units Sold", "Revenue"]
    assert rows[1] == ["Widget A", "150", "4500"]
    assert rows[2] == ["Widget B", "80", "2400"]


def test_parse_cell_value_converts_numeric_types():
    """Verify numeric strings become ints/floats without stripping codes."""
    assert _parse_cell_value("123") == 123
    assert isinstance(_parse_cell_value("123"), int)

    assert _parse_cell_value("45.67") == 45.67
    assert isinstance(_parse_cell_value("45.67"), float)

    assert _parse_cell_value("$1,250") == 1250
    assert _parse_cell_value("true") is True
    assert _parse_cell_value("false") is False

    # Preserves leading zeros on codes and string IDs
    assert _parse_cell_value("01234") == "01234"
    assert _parse_cell_value("STU-001") == "STU-001"
    assert _parse_cell_value("John Doe") == "John Doe"


def test_excel_generation_numeric_cells_and_types():
    """Verify generated XLSX contains real numeric cells, not raw strings."""
    content = (
        "Item,Quantity,Unit Price,In Stock\n"
        "Laptop,15,1200.50,true\n"
        "Monitor,40,250,false\n"
    )
    test_path = OUTPUT_DIR / "test_numeric_types.xlsx"
    _generate_xlsx("Inventory Report", content, test_path)

    wb = openpyxl.load_workbook(test_path)
    ws = wb.active

    # Row 1 is Title bar
    assert ws.cell(1, 1).value.startswith("Inventory Report")

    # Row 2 is Header
    assert ws.cell(2, 1).value == "Item"
    assert ws.cell(2, 2).value == "Quantity"

    # Row 3 is Data
    assert ws.cell(3, 1).value == "Laptop"
    assert ws.cell(3, 2).value == 15
    assert isinstance(ws.cell(3, 2).value, int)
    assert ws.cell(3, 3).value == 1200.50
    assert isinstance(ws.cell(3, 3).value, float)
    assert ws.cell(3, 4).value is True

    # Font should be Calibri
    assert ws.cell(2, 1).font.name == "Calibri"
    assert ws.cell(3, 1).font.name == "Calibri"


def test_tools_generate_xlsx_numeric_types():
    """Verify tools.generate_xlsx also parses numbers into ints/floats."""
    headers = ["ID", "Score", "Active"]
    data = [["STU-1", "98", "true"], ["STU-2", "85.5", "false"]]
    res = tools.generate_xlsx("Student Grades", headers, data, "pytest_tools_grades.xlsx")
    assert "saved to:" in res

    wb = openpyxl.load_workbook(OUTPUT_DIR / "pytest_tools_grades.xlsx")
    ws = wb.active
    assert ws.cell(2, 2).value == 98
    assert isinstance(ws.cell(2, 2).value, int)
    assert ws.cell(3, 2).value == 85.5
    assert isinstance(ws.cell(3, 2).value, float)


def test_word_unclosed_code_block_flushes():
    """Verify that unclosed code blocks at EOF are rendered properly in Word."""
    content = (
        "## Python Implementation\n\n"
        "```python\n"
        "def greet(name):\n"
        "    return f'Hello, {name}'"
        # No closing ```
    )
    artifact = create_artifact("Code Guide", content, "docx", "pytest_unclosed_code.docx")
    doc = Document(str(artifact.path))

    # Should contain the code text in a shaded table
    all_table_text = " ".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "def greet(name):" in all_table_text


def test_tools_write_docx_parity_with_tables_and_callouts():
    """Verify tools.write_docx parses tables and callouts without dumping raw markdown."""
    content = (
        "## Performance Metrics\n\n"
        "| Metric | Value |\n"
        "| Latency | 12ms |\n"
        "| Throughput | 5000 rps |\n\n"
        "> Note: These metrics were measured under load.\n"
    )
    res = tools.write_docx("System Performance", content, "pytest_tools_performance.docx")
    assert "saved to:" in res

    doc = Document(str(OUTPUT_DIR / "pytest_tools_performance.docx"))
    # Table should be created (masthead + data table)
    assert len(doc.tables) >= 2
    assert len(doc.inline_shapes) == 0  # No logo image


def test_multi_agent_splits_sentences_with_periods():
    """Verify multi-action queries separated by periods split into self-contained sub-tasks."""
    query = "What is python. Give me an essay on harshad mehta. Give me a word file showing investment principle"
    parts = router._split_multi_actions(query)
    assert len(parts) == 3
    assert "What is python" in parts[0]
    assert "essay on harshad mehta" in parts[1].lower()
    assert "word file" in parts[2].lower()


def test_multi_agent_with_file_subtask_does_not_hijack_entire_query():
    """Verify that file intent in a multi-task query flags is_multi_part: True rather than single-file hijack."""
    query = "What is python. Give me an essay on harshad mehta. Give me a word file showing investment principle"
    info = router.classify_question(query)
    assert info["is_multi_part"] is True
    assert info["category"] == "agent_task"

    tasks = router.break_into_tasks(query)
    assert len(tasks) >= 3


def test_word_title_cleaning_conversational_lead_in():
    """Verify that conversational intention lead-ins are stripped from Word document titles."""
    from artifacts import _clean_topic_from_question, derive_clean_title
    query = "I want to watch best tmkoc episodes ever. so give me the word file showing all the best episodes of tmkoc ever"
    clean_topic = _clean_topic_from_question(query)
    assert "Best" in clean_topic and "Episodes" in clean_topic and "TMKOC" in clean_topic.upper()
    assert "want to watch" not in clean_topic.lower()
    assert "so give me" not in clean_topic.lower()

    title, _, slug = derive_clean_title(query, "Content about episodes")
    assert "want to" not in title.lower()
    assert "best_tmkoc_episodes" in slug or "episodes" in slug


def test_langgraph_clean_subject_conversational():
    """Verify _extract_clean_subject and _file_topic strip conversational wrappers."""
    query = "I want to watch best tmkoc episodes ever. so give me the word file showing all the best episodes of tmkoc ever"
    subj = langgraph_agent._extract_clean_subject(query, "docx")
    assert "want to watch" not in subj.lower()
    assert "so give me" not in subj.lower()
    assert "episodes" in subj.lower()


def test_word_inline_code_formatting():
    """Verify inline `code` backticks format cleanly in DOCX with Consolas font."""
    content = (
        "## Configuration Guide\n\n"
        "Set the variable `MAX_TURNS = 10` and call `clear_history()` to reset.\n"
    )
    artifact = create_artifact("Config Guide", content, "docx", "pytest_inline_code.docx")
    doc = Document(str(artifact.path))
    # Verify paragraphs exist and font settings applied
    has_consolas = False
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.name == "Consolas":
                has_consolas = True
                break
    assert has_consolas is True

