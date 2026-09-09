from pathlib import Path

import db
import langgraph_agent
from artifacts import OUTPUT_DIR, create_artifact
from fastapi.testclient import TestClient
from main import app


def test_supported_artifacts_are_created_and_readable():
    for extension in ("txt", "md", "json", "csv", "xlsx", "docx", "pdf"):
        artifact = create_artifact("Test Artifact", "name,value\nagent,1", extension, f"pytest_artifact.{extension}")
        assert Path(artifact.path).is_file()
        assert artifact.size_bytes > 0


def test_agent_completes_requested_file_workflow(monkeypatch):
    monkeypatch.setattr(langgraph_agent, "_llm_content", lambda *_: "def last_non_zero(n):\n    return n")
    result = langgraph_agent.run_agent("Generate Python code for the last non-zero factorial digit and provide it as a PDF.")
    assert result["artifact"] is not None
    assert result["artifact"]["filename"].endswith(".pdf")
    assert any(event["stage"] == "creating_file" for event in result["events"])
    assert (OUTPUT_DIR / result["artifact"]["filename"]).is_file()


def test_sqlite_session_persistence():
    assert db.init_db()
    session = "pytest-offline-session"
    assert db.create_session(session)
    assert db.save_message(session, "user", "remember this")
    messages = db.get_session_messages(session)
    assert messages[-1]["content"] == "remember this"
    assert db.delete_session(session)


def test_agent_api_returns_download_url(monkeypatch):
    monkeypatch.setattr(langgraph_agent, "_llm_content", lambda *_: "A concise report.")
    with TestClient(app) as client:
        response = client.post("/ask/agent", json={"query": "Create a PDF about offline AI."})
        artifact = response.json()["artifact"]
        download = client.get(artifact["download_url"])
    assert response.status_code == 200
    assert artifact["download_url"].startswith("/files/")
    assert download.status_code == 200
    assert download.headers["content-type"] == "application/pdf"


def test_capabilities_route_is_available():
    with TestClient(app) as client:
        capabilities = client.get("/capabilities")
    assert capabilities.json()["offline"] is True


def test_derive_clean_title_removes_prompt_instructions():
    from artifacts import derive_clean_title

    # 1. Title extracted from markdown heading inside content
    q1 = "Give me the code to print the fibonaaci series in python and also generatet he pdf"
    c1 = "# Python code to print the Fibonacci series\n\ndef fib(n):\n    pass"
    title1, content1, slug1 = derive_clean_title(q1, c1)
    assert title1 == "Python code to print the Fibonacci series"
    assert not content1.startswith("# Python code")
    assert not title1.lower().startswith("give me")

    # 2. Title cleaned directly from prompt without content heading
    q2 = "Give me the code to print numebrs from 1 to 10 in python also generate the pdf"
    c2 = "for i in range(1, 11):\n    print(i)"
    title2, content2, slug2 = derive_clean_title(q2, c2)
    assert not title2.lower().startswith("give me")
    assert "generate" not in title2.lower()
    assert "pdf" not in title2.lower()
    assert "1 to 10" in title2


def test_clean_title_uses_subject_when_model_omits_heading():
    from artifacts import derive_clean_title

    title, content, slug = derive_clean_title(
        "Create a PowerPoint presentation about solar energy adoption trends", "Plain body text."
    )
    assert "presentation" not in title.lower()
    assert title.lower() == "solar energy adoption trends"
    assert content == "Plain body text."
    assert slug == "solar_energy_adoption_trends"


def test_student_dataset_honors_requested_record_count():
    from langgraph_agent import _student_dataset_csv

    dataset = _student_dataset_csv("Generate a dataset of 50 students in Excel")
    assert dataset is not None
    assert len(dataset.splitlines()) == 51  # header + exactly 50 student records


def test_partial_csv_is_extended_to_requested_row_count():
    from langgraph_agent import _normalise_csv_row_count

    csv_text = "ID,Name\n1,Ada\n2,Grace"
    assert len(_normalise_csv_row_count(csv_text, 50).splitlines()) == 51


def test_csv_honors_requested_column_count():
    from langgraph_agent import _normalise_csv_row_count

    rows = _normalise_csv_row_count("ID,Name\n1,Ada", 3, 5).splitlines()
    assert len(rows) == 4
    assert all(len(line.split(",")) == 5 for line in rows)


def test_presentation_source_is_extended_to_requested_slide_count():
    from langgraph_agent import _ensure_presentation_slide_count

    content = "# Introduction\n- First point\n- Second point"
    deck_source = _ensure_presentation_slide_count(content, "Renewable Energy", 12)
    assert len([line for line in deck_source.splitlines() if line.startswith("# ")]) == 12


def test_generic_slide_labels_are_replaced_with_professional_titles():
    from langgraph_agent import _professionalise_slide_titles

    result = _professionalise_slide_titles("# Slide_1\n- Point\n# Slide 2\n- Point", "Renewable Energy")
    assert "Slide_1" not in result
    assert "Slide 2" not in result
    assert "Renewable Energy — Executive Summary" in result


def test_instruction_heading_does_not_become_document_title():
    from artifacts import derive_clean_title

    title, _, _ = derive_clean_title(
        "Create a PDF about renewable energy adoption",
        "# Generate a PDF about renewable energy adoption\n\nUseful content.",
    )
    assert title.lower() == "renewable energy adoption"


def test_is_file_creation_request_detection():
    from ask import is_file_creation_request

    assert is_file_creation_request("Give me the code to print numebrs from 1 to 10 in python also generate the pdf") is True
    assert is_file_creation_request("Give me the code to print the fibonaaci series in python and also generatet he pdf") is True
    assert is_file_creation_request("Create a word document with project plan") is True
    assert is_file_creation_request("Export sales data to excel") is True
    assert is_file_creation_request("How do I reverse a string in python?") is False


def test_router_find_embedded_tool_calls():
    from router import _find_embedded_tool_calls

    model_output = (
        "Sure, here is the code:\n```python\nprint(1)\n```\n"
        "To generate a PDF from this text, use:\n"
        "```json\n"
        '{"name": "generate_pdf_from_text", "arguments": {"title": "Numbers", "content": "1", "filename": "numbers.pdf"}}\n'
        "```"
    )
    tools = _find_embedded_tool_calls(model_output)
    assert len(tools) == 1
    assert tools[0][0] == "generate_pdf_from_text"
    assert tools[0][1]["title"] == "Numbers"


def test_universal_word_prompt_understanding():
    from langgraph_agent import detect_file_intent
    from ask import is_file_creation_request
    from router import classify_question

    prompts = [
        "make a word file about sales strategy",
        "create doc file on cloud architecture",
        "write python script for fibonacci and generate a word file for it",
        "draft a docx report on energy efficiency",
        "prepare ms word file with project roadmap",
        "generate word document covering market trends",
        "convert to docx format",
        "give me a word doc for data analysis",
    ]

    for p in prompts:
        intent = detect_file_intent(p)
        assert intent["file_format"] == "docx", f"Failed file intent for prompt: '{p}'"
        assert is_file_creation_request(p) is True, f"Failed creation check for prompt: '{p}'"
        cls = classify_question(p)
        assert cls["category"] == "agent_task", f"Failed router classification for prompt: '{p}'"


def test_docx_artifact_creation_with_format_aliases():
    for fmt in ("docx", "doc", "word", "msword"):
        artifact = create_artifact("Project Report", "# Heading 1\n\n- Bullet 1\n- Bullet 2", fmt, f"test_alias_{fmt}.docx")
        assert Path(artifact.path).is_file()
        assert artifact.size_bytes > 0
        assert artifact.path.name.endswith(".docx")


def test_agent_completes_word_file_workflow(monkeypatch):
    monkeypatch.setattr(langgraph_agent, "_llm_content", lambda *_: "# Market Analysis\n\nKey trends in 2026.")
    result = langgraph_agent.run_agent("Write a python code for data processing and put it in a word file.")
    assert result["artifact"] is not None
    assert result["artifact"]["filename"].endswith(".docx")
    assert (OUTPUT_DIR / result["artifact"]["filename"]).is_file()


def test_introduction_not_used_as_document_title():
    from artifacts import derive_clean_title

    q = "give me a report on Artificial Intelligence in docx and pdf and ppt"
    c = "## Introduction\nArtificial Intelligence is transforming industries.\n## Core Concepts\nNeural networks."
    title, content, slug = derive_clean_title(q, c)

    assert title == "Artificial Intelligence"
    assert slug == "artificial_intelligence"
    assert "## Introduction" in content

    # Even with a single '#' introduction
    c2 = "# Introduction\nOverview text.\n## Details\nSome details."
    title2, content2, slug2 = derive_clean_title(q, c2)
    assert title2 == "Artificial Intelligence"
    assert slug2 == "artificial_intelligence"
    assert "# Introduction" in content2 or "## Introduction" in content2


def test_word_document_has_no_logo_and_has_branding():
    from docx import Document

    content = (
        "## Introduction\n"
        "AI systems are advancing rapidly.\n\n"
        "### Key Principles\n"
        "- Scalability\n"
        "- Interpretability\n\n"
        "| Architecture | Latency |\n"
        "| Transformer  | Low     |\n"
    )
    artifact = create_artifact("Cybersecurity Best Practices", content, "docx", "test_structure.docx")
    doc = Document(str(artifact.path))

    # 1. NO logo image in the document
    assert len(doc.inline_shapes) == 0, "Word document should not contain logo picture"

    # 2. Branding text is present
    all_text = " ".join(p.text for p in doc.paragraphs)
    table_text = " ".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    combined = f"{all_text} {table_text}"
    assert "AGENT OTG" in combined
    assert "DWE TEAM" in combined

    # 3. Main heading is the topic, NOT 'Introduction'
    title_paragraphs = [p for p in doc.paragraphs if "Cybersecurity Best Practices" in p.text]
    assert len(title_paragraphs) >= 1
    # Check that 'Introduction' is not the cover title
    assert doc.paragraphs[1].text != "Introduction"

    # 4. Content is structured: table has columns formatted, bullet points present
    assert len(doc.tables) >= 2  # masthead table + markdown data table
    data_table = doc.tables[1]
    assert len(data_table.columns) == 2
    assert data_table.columns[0].width is not None


